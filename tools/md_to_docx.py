"""把技术文档 Markdown 转成排版良好的 Word 文档。

为什么不用 pandoc：本机没装，而且中文字体需要显式设 `w:eastAsia`，
代码块需要等宽字体 + 底纹，表格需要真表格——用 python-docx 直接控制更稳。

用法:
    python tools/md_to_docx.py <input.md> <output.docx> [标题] [副标题]
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CODE_BG = "F2F4F7"
QUOTE_BG = "FFF8E1"


def set_cjk(run, font_name: str | None = None, size: int | None = None,
            bold: bool | None = None, color=None, mono: bool = False) -> None:
    """设置含中文的 run 字体（必须显式写 w:eastAsia，否则中文会掉回默认字体）。

    Args:
        run: python-docx 的 run。
        font_name: 字体名。
        size: 字号（pt）。
        bold: 是否加粗。
        color: 颜色。
        mono: 是否用等宽字体。
    """
    name = CODE_FONT if mono else (font_name or BODY_FONT)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph, fill: str) -> None:
    """给段落加底纹。

    Args:
        paragraph: 段落。
        fill: 十六进制颜色（不带 #）。
    """
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def left_bar(paragraph, color: str = "F0B429") -> None:
    """给段落加左侧竖条（用于引用块）。

    Args:
        paragraph: 段落。
        color: 十六进制颜色。
    """
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    borders.append(left)
    ppr.append(borders)


BOLD = re.compile(r"\*\*(.+?)\*\*")
CODE = re.compile(r"`([^`]+)`")


def scan_inline(text: str):
    """把一行文本切成 (文本, 加粗, 等宽) 三元组序列。

    为什么不用单个大正则做 `split`：那样 `**…`code`…**` 这种"加粗套代码"会被
    第一个分支整个吃掉，里面的反引号原样留在正文里（第一版就是这么错的）。
    这里按"先找最靠前的标记，再递归处理它的内容"来做，天然支持嵌套。

    Args:
        text: 含 `**加粗**` 与 `` `代码` `` 标记的文本。

    Yields:
        (片段文本, 是否加粗, 是否等宽)。
    """
    pos = 0
    while pos < len(text):
        mb = BOLD.search(text, pos)
        mc = CODE.search(text, pos)
        # 取最靠前的标记；同位置时加粗优先（它是外层）
        nxt = None
        if mb and mc:
            nxt = mb if mb.start() <= mc.start() else mc
        else:
            nxt = mb or mc
        if nxt is None:
            yield text[pos:], False, False
            return
        if nxt.start() > pos:
            yield text[pos:nxt.start()], False, False
        if nxt is mb:
            # 加粗：内容再扫一遍，于是里面的 `代码` 也能正确变等宽
            for sub, _bold, mono in scan_inline(mb.group(1)):
                yield sub, True, mono
        else:
            yield mc.group(1), False, True
        pos = nxt.end()


def add_inline(paragraph, text: str, size: int = 10.5, base_bold: bool = False) -> None:
    """把一行文本按 **加粗** / `代码` 拆成多个 run（支持嵌套）。

    Args:
        paragraph: 目标段落。
        text: 含轻量标记的文本。
        size: 字号。
        base_bold: 整段是否加粗。
    """
    for piece, bold, mono in scan_inline(text):
        if not piece:
            continue
        run = paragraph.add_run(piece)
        set_cjk(
            run,
            size=size - 0.5 if mono else size,
            bold=bold or base_bold or None,
            mono=mono,
            color=RGBColor(0xB0, 0x30, 0x60) if mono else None,
        )


def add_md_paragraph(doc, text: str, style: str | None = None,
                     indent: float | None = None, size: float = 10.5,
                     quote: bool = False):
    """新建一个段落并写入 markdown 文本。

    Args:
        doc: 文档。
        text: 文本。
        style: 段落样式名。
        indent: 左缩进（pt）。
        size: 字号。
        quote: 是否按引用块样式渲染。

    Returns:
        新建的段落。
    """
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if indent is not None:
        para.paragraph_format.left_indent = Pt(indent)
    if quote:
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        shade(para, QUOTE_BG)
        left_bar(para)
    add_inline(para, text, size=size)
    return para


def strip_md_links(text: str) -> str:
    """把 markdown 链接降级成纯文字。

    Word 里保留 `#anchor` 没有意义（页内锚点不通用），外部 URL 又会让正文变长，
    所以统一只留链接文字。注意：**只去掉链接的 URL 部分，正文里的 `**加粗**`
    与 `` `代码` `` 一律原样保留**，交给 `scan_inline` 处理——早期版本在这里
    做了字符串拼装，结果把标记弄坏过。

    Args:
        text: 原文本。

    Returns:
        处理后的文本。
    """
    return re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)


def build(md_path: Path, out_path: Path, title: str, subtitle: str) -> int:
    """执行转换。

    Args:
        md_path: 输入 Markdown。
        out_path: 输出 docx。
        title: 封面标题。
        subtitle: 封面副标题。

    Returns:
        写入的块数。
    """
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    # 页边距
    for section in doc.sections:
        section.left_margin = Pt(64)
        section.right_margin = Pt(64)
        section.top_margin = Pt(64)
        section.bottom_margin = Pt(64)

    # ---------------- 封面 ----------------
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cjk(p.add_run(title), size=26, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cjk(p.add_run(subtitle), size=13, color=RGBColor(0x55, 0x55, 0x55))
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cjk(p.add_run("⚠ 文中所有实测结论均标注了取证方式与日期；"
                      "标为「未证实」的假设请勿当作结论引用。"),
            size=9.5, color=RGBColor(0x88, 0x44, 0x00))
    doc.add_page_break()

    blocks = 0
    i = 0
    in_code = False
    code_buf: list[str] = []

    def flush_code() -> None:
        """把累积的代码行写成一个底纹段落。"""
        nonlocal code_buf, blocks
        if not code_buf:
            return
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.left_indent = Pt(10)
        shade(para, CODE_BG)
        for idx, cline in enumerate(code_buf):
            if idx:
                para.add_run().add_break()
            run = para.add_run(cline if cline else " ")
            set_cjk(run, size=9, mono=True, color=RGBColor(0x1A, 0x1A, 0x1A))
        code_buf = []
        blocks += 1

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # 代码围栏
        if line.strip().startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        # 表格（连续的 | 行）
        if line.lstrip().startswith("|") and line.rstrip().endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                if re.fullmatch(r"\|[\s:|-]+\|", tl):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncol)
                table.style = "Light Grid Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ridx, cells in enumerate(rows):
                    row = table.add_row()
                    for cidx in range(ncol):
                        text = cells[cidx] if cidx < len(cells) else ""
                        cell = row.cells[cidx]
                        cell.text = ""
                        para = cell.paragraphs[0]
                        add_inline(
                            para,
                            strip_md_links(text),
                            size=9,
                            base_bold=(ridx == 0),
                        )
                doc.add_paragraph()
                blocks += 1
            continue

        stripped = line.strip()

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = strip_md_links(m.group(2))
            if level == 1:
                if blocks:  # 正文标题另起一页
                    doc.add_page_break()
                para = add_md_paragraph(doc, text, size=19)
                for run in para.runs:
                    run.font.color.rgb = ACCENT
                para.paragraph_format.space_after = Pt(10)
            else:
                size = {2: 15, 3: 12.5}.get(level, 11)
                para = add_md_paragraph(doc, text, size=size)
                for run in para.runs:
                    run.font.bold = True
                    if level == 2:
                        run.font.color.rgb = RGBColor(0x2E, 0x5C, 0x8A)
                    elif level == 3:
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                para.paragraph_format.space_before = Pt(14 if level == 2 else 10)
                para.paragraph_format.space_after = Pt(6 if level == 2 else 4)
            blocks += 1
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            merged = strip_md_links(" ".join(x for x in buf if x))
            add_md_paragraph(doc, merged, indent=12, size=9.5, quote=True)
            blocks += 1
            continue

        # 列表
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", raw)
        if m:
            indent, _, content = m.groups()
            add_md_paragraph(
                doc,
                strip_md_links(content),
                style="List Bullet",
                indent=18 + 14 * (len(indent) // 2),
            )
            blocks += 1
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        add_md_paragraph(doc, strip_md_links(stripped))
        blocks += 1
        i += 1

    flush_code()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return blocks


def main() -> int:
    """入口。

    Returns:
        退出码。
    """
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    md = Path(sys.argv[1])
    out = Path(sys.argv[2])
    title = sys.argv[3] if len(sys.argv) > 3 else "话题路由器 PRO · 技术文档"
    subtitle = sys.argv[4] if len(sys.argv) > 4 else ""
    n = build(md, out, title, subtitle)
    print(f"已写出 {out}（{out.stat().st_size} 字节，共 {n} 个内容块）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
